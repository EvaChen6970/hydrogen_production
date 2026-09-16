"""Build an English Word report from the completed model prediction files."""

import json
import os
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "outputs")
FIG = os.path.join(OUT, "figures")


def metrics(pred, true):
    pred, true = np.asarray(pred), np.asarray(true)
    e = pred - true
    mask = np.abs(true) > 1e-6
    ss = np.sum((true - true.mean()) ** 2)
    return {"MAE": float(np.mean(np.abs(e))), "RMSE": float(np.sqrt(np.mean(e ** 2))),
            "MAPE": float(np.mean(np.abs(e[mask] / true[mask])) * 100),
            "R2": float(1 - np.sum(e ** 2) / ss) if ss else float("nan")}


FILES = {
    "HydroFormerLSTM": "hydroformer_test_predictions.csv",
    "XGBoost": "xgboost_test_predictions.csv",
    "1D-CNN": "cnn_test_predictions.csv",
    "GRU": "gru_test_predictions.csv",
    "LSTM": "lstm_test_predictions.csv",
    "Transformer": "transformer_test_predictions.csv",
    "PT-MELT": "ptmelt_aligned_predictions.csv",
}


def load_predictions():
    data = {}
    for name, file_name in FILES.items():
        path = os.path.join(OUT, file_name)
        if os.path.exists(path):
            df = pd.read_csv(path)
            if {"experiment", "experiment_position"}.issubset(df.columns):
                df["experiment"] = df["experiment"].astype(str)
                df["experiment_position"] = pd.to_numeric(df["experiment_position"], errors="coerce")
                df["_comparison_key"] = df["experiment"] + "|" + df["experiment_position"].astype("Int64").astype(str)
            data[name] = df
    keyed = [df for df in data.values() if "_comparison_key" in df.columns]
    if keyed:
        common = set(keyed[0]["_comparison_key"])
        for df in keyed[1:]:
            common &= set(df["_comparison_key"])
        for name, df in list(data.items()):
            if "_comparison_key" in df.columns:
                data[name] = df[df["_comparison_key"].isin(common)].copy()
                data[name].drop(columns=["_comparison_key"], inplace=True)
        print(f"common_comparison_samples={len(common)}")
    return data


def make_figures(data):
    os.makedirs(FIG, exist_ok=True)
    names, ms = [], []
    for name, df in data.items():
        names.append(name)
        ms.append(metrics(df.h2_pred, df.h2_true))
    fig, ax = plt.subplots(figsize=(11, 5), dpi=160)
    vals = [m["MAE"] for m in ms]
    bars = ax.bar(names, vals, color=["#C44E52", "#4C72B0", "#55A868", "#8172B2", "#DD8452", "#937860", "#0072B2"])
    ax.set_ylabel("H2 MAE (kg/h)")
    ax.set_title("Hydrogen-production MAE across evaluated models")
    ax.grid(axis="y", alpha=.25)
    ax.tick_params(axis="x", rotation=25)
    for bar, value in zip(bars, vals):
        ax.text(bar.get_x() + bar.get_width()/2, value, f"{value:.2f}", ha="center", va="bottom", fontsize=8)
    fig.tight_layout()
    comparison = os.path.join(FIG, "english_report_model_comparison.png")
    fig.savefig(comparison, bbox_inches="tight")
    plt.close(fig)

    selected = [n for n in ["PT-MELT", "HydroFormerLSTM", "LSTM", "GRU", "Transformer"] if n in data]
    n = min(1000, min(len(data[x]) for x in selected))
    fig, axes = plt.subplots(len(selected), 1, figsize=(12, 2.2 * len(selected)), sharex=True, dpi=150)
    if len(selected) == 1:
        axes = [axes]
    for ax, name in zip(axes, selected):
        df = data[name]
        ax.plot(df.h2_true.iloc[:n].to_numpy(), color="black", lw=1.0, label="Ground truth")
        ax.plot(df.h2_pred.iloc[:n].to_numpy(), lw=.9, label=name)
        ax.set_ylabel("kg/h")
        ax.set_title(name)
        ax.grid(alpha=.2)
        ax.legend(loc="upper right", fontsize=8)
    axes[-1].set_xlabel("Prediction sample")
    fig.suptitle("Hydrogen-production predictions on each model output file", y=.995)
    fig.tight_layout()
    curves = os.path.join(FIG, "english_report_prediction_curves.png")
    fig.savefig(curves, bbox_inches="tight")
    plt.close(fig)
    return comparison, curves


def add_table(doc, rows, headers):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = str(text)
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = str(text)
    return table


def main():
    data = load_predictions()
    comparison_fig, curves_fig = make_figures(data)
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    title = doc.add_heading("Hydrogen Production Forecasting: Model Comparison Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Experimental report generated from the NLR simulated-wind PEM electrolysis dataset.")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("1. Dataset and Task", level=1)
    doc.add_paragraph(
        "The experiment uses the National Laboratory of the Rockies (NLR) simulated wind electrolysis dataset. "
        "The release contains 61,285 samples at 1 Hz from 13 experiments: multiple wind speeds, turbulence classes, "
        "electrolyzer configurations, current ramp rates, and a steady-state characterization profile. The measured "
        "target is hydrogen flow, IVAL_f_FM011_Flow, in kg/h. LCOH is a derived auxiliary target and is not a direct sensor measurement."
    )
    add_table(doc, [["Total rows", "61,285"], ["Sampling rate", "1 Hz"], ["Experiments", "13"],
                    ["Input window", "180 seconds"], ["Forecast horizon", "30 seconds"], ["Window stride", "15 seconds"],
                    ["Train / validation / test windows", "909 / 681 / 681"]], ["Item", "Value"])

    doc.add_heading("2. Data Sources and Links", level=1)
    doc.add_paragraph("Dataset: NLR Public Reference Data for Megawatt-Scale Hydrogen Electrolysis – Simulated Wind")
    doc.add_paragraph("Official dataset page: https://data.nlr.gov/submissions/305")
    doc.add_paragraph("PT-MELT model: NLR published hydrogen electrolysis LSTM model")
    doc.add_paragraph("Hugging Face model repository: https://huggingface.co/NatLabRockies/ptmelt-hydrogen-electrolysis")
    doc.add_paragraph("The dataset page provides the combined CSV file and the individual wind and characterization experiment archives. The PT-MELT repository provides the published model weights, metadata, and normalization parameters used for the PT-MELT baseline.")

    doc.add_heading("3. Data Processing and Split", level=1)
    doc.add_paragraph(
        "Raw columns were converted to numeric values, missing values were forward-filled and linearly interpolated, "
        "and features were standardized using training data. Wind speed and turbulence were derived from the wind-power signal. "
        "The split is phase-balanced by experiment: phases 0, 3, 6, and 9 for training; phases 1, 4, and 7 for validation; "
        "and phases 2, 5, and 8 for testing. Windows never cross experiment or phase boundaries."
    )

    doc.add_heading("4. Model Settings", level=1)
    add_table(doc, [
        ["HydroFormerLSTM", "Conv1D + 4-layer Transformer + 2-layer LSTM; d_model=192; 8 heads; LSTM hidden=256; 5.29M parameters"],
        ["Transformer", "4-layer Transformer; d_model=128; 8 heads; feed-forward width=768"],
        ["LSTM", "2-layer LSTM; hidden size=128"],
        ["GRU", "2-layer GRU; hidden size=128"],
        ["1D-CNN", "Four dilated residual convolution blocks; 64 channels"],
        ["XGBoost", "300 trees; max depth=8; learning rate=0.05; histogram tree method"],
        ["PT-MELT", "Published NLR single-layer attention LSTM; 15 raw electrolysis sensors; sequence length=60; one-step model evaluated at 30 seconds"],
    ], ["Model", "Configuration"])

    doc.add_heading("5. Results", level=1)
    rows = []
    for name, df in data.items():
        m = metrics(df.h2_pred, df.h2_true)
        rows.append([name, len(df), f"{m['MAE']:.4f}", f"{m['RMSE']:.4f}", f"{m['MAPE']:.2f}%", f"{m['R2']:.4f}"])
    add_table(doc, rows, ["Model", "Samples", "H2 MAE", "H2 RMSE", "H2 MAPE", "H2 R2"])
    doc.add_paragraph(
        "To make the comparison fair, every model is evaluated on the exact common intersection of experiment and "
        "experiment_position keys. PT-MELT is first aligned to the HydroFormerLSTM test windows, and the report then "
        "filters all prediction files to the same shared keys. Consequently, the sample count and ground-truth sequence "
        "are identical across all models in the table and figures."
    )
    doc.add_picture(comparison_fig, width=Inches(6.7))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(curves_fig, width=Inches(6.7))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("6. Interpretation", level=1)
    doc.add_paragraph(
        "The improved HydroFormerLSTM is substantially better than the original Transformer and 1D-CNN on this dataset. "
        "However, the compact recurrent baselines remain competitive: LSTM and GRU exploit the short, strongly autocorrelated "
        "electrolyzer dynamics efficiently. This indicates that increasing parameter count alone is not sufficient for this sample size. "
        "PT-MELT provides a strong published reference model and confirms that the 15 raw electrolysis sensors are highly predictive. "
        "MAPE should be interpreted cautiously because near-zero hydrogen-flow values inflate percentage errors."
    )
    doc.add_heading("7. Reproducibility Commands", level=1)
    doc.add_paragraph("python main.py --model hydroformer --epochs 50\npython scripts/evaluate_ptmelt.py --horizon 30\npython scripts/compare_predictions.py")
    path = os.path.join(OUT, "hydrogen_model_comparison_report.docx")
    doc.save(path)
    print(path)


if __name__ == "__main__":
    main()
